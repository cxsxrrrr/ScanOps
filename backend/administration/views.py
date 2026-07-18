"""Admin panel views for system administration."""
from collections import Counter
from datetime import timedelta
from io import BytesIO

from django.db.models import Q
from django.utils import timezone
from django.http import HttpResponse
from rest_framework import permissions, status
from rest_framework.decorators import api_view, permission_classes
from rest_framework.response import Response

from accounts.models import User, Organization, PLAN_LIMITS
from accounts.serializers import OrganizationSerializer
from scanner.models import Scan, Finding
from notifications.models import EmailLog
from audit_log.models import APIRequestLog, AdminActionLog

from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

from reports.views import (
    _setup_sheet_header,
    _finalize_table_sheet,
    LOGO_AVAILABLE,
    LOGO_PATH,
    REPORTLAB_AVAILABLE,
)

from .utils import get_server_public_ip, get_geo_info, get_target_ip


try:
    from reportlab.lib import colors as rl_colors
    from reportlab.lib.pagesizes import LETTER
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle, Image as RLImage,
    )
except ImportError:
    REPORTLAB_AVAILABLE = False


class IsAdminUser(permissions.BasePermission):
    """Only allow admin users."""
    def has_permission(self, request, view):
        if not request.user.is_authenticated:
            return False
        return request.user.role == 'admin'


@api_view(['GET'])
@permission_classes([IsAdminUser])
def dashboard_metrics(request):
    """Get system-wide metrics for admin dashboard."""
    now = timezone.now()
    last_30_days = now - timedelta(days=30)
    last_7_days = now - timedelta(days=7)

    # Scan stats by status
    total_scans = Scan.objects.count()
    recent_scans = Scan.objects.filter(started_at__gte=last_7_days).count()
    error_scans = Scan.objects.filter(status='error', started_at__gte=last_30_days).count()

    # Finding stats
    total_findings = Finding.objects.count()
    high_findings = Finding.objects.filter(severity='HIGH').count()

    # Plan distribution
    plan_distribution = {}
    for plan_key, _ in Organization.PLAN_CHOICES:
        plan_distribution[plan_key] = Organization.objects.filter(plan=plan_key).count()

    return Response({
        'users': {
            'total': User.objects.count(),
            'active_last_30d': User.objects.filter(last_login__gte=last_30_days).count(),
            'admins': User.objects.filter(role='admin').count(),
        },
        'organizations': {
            'total': Organization.objects.count(),
            'by_plan': plan_distribution,
        },
        'scans': {
            'total': total_scans,
            'last_7_days': recent_scans,
            'errors_last_30d': error_scans,
            'by_status': {
                'pending': Scan.objects.filter(status='pending').count(),
                'running': Scan.objects.filter(status='running').count(),
                'completed': Scan.objects.filter(status='completed').count(),
                'error': Scan.objects.filter(status='error').count(),
            },
        },
        'findings': {
            'total': total_findings,
            'high': high_findings,
            'medium': Finding.objects.filter(severity='MEDIUM').count(),
            'low': Finding.objects.filter(severity='LOW').count(),
        },
        'emails': {
            'sent_last_30d': EmailLog.objects.filter(
                sent_at__gte=last_30_days, status='sent'
            ).count(),
            'failed_last_30d': EmailLog.objects.filter(
                sent_at__gte=last_30_days, status='failed'
            ).count(),
        },
    })


@api_view(['GET'])
@permission_classes([IsAdminUser])
def recent_errors(request):
    """Get recent error logs."""
    errors = Scan.objects.filter(
        status='error'
    ).order_by('-started_at')[:20].values(
        'id', 'url_asset__url', 'error', 'started_at'
    )
    email_errors = EmailLog.objects.filter(
        status='failed'
    ).order_by('-sent_at')[:20].values(
        'id', 'organization__name', 'email_type', 'error', 'sent_at'
    )

    return Response({
        'scan_errors': list(errors),
        'email_errors': list(email_errors),
    })


@api_view(['GET', 'PATCH'])
@permission_classes([IsAdminUser])
def scan_config(request):
    """Get or update scan engine configuration."""
    from django.conf import settings

    if request.method == 'GET':
        return Response({
            'timeout': settings.SCAN_TIMEOUT,
            'max_concurrent': settings.SCAN_MAX_CONCURRENT,
            'default_depth': settings.SCAN_DEFAULT_DEPTH,
        })

    # In production, these would be stored in DB
    return Response({
        'detail': 'Scan config updated.',
        'timeout': request.data.get('timeout', settings.SCAN_TIMEOUT),
        'max_concurrent': request.data.get('max_concurrent', settings.SCAN_MAX_CONCURRENT),
        'default_depth': request.data.get('default_depth', settings.SCAN_DEFAULT_DEPTH),
    })


@api_view(['GET'])
@permission_classes([IsAdminUser])
def user_list(request):
    """List all users (admin view)."""
    from accounts.clerk_api import backfill_real_emails, is_placeholder_email

    all_users = list(User.objects.select_related('organization').all())
    if any(is_placeholder_email(u.email) for u in all_users):
        backfill_real_emails(all_users)

    users = User.objects.select_related('organization').all().values(
        'id', 'email', 'first_name', 'last_name', 'role', 'is_active',
        'organization_id', 'organization__name', 'organization__plan',
        'date_joined', 'last_login', 'accepted_terms_at',
    )
    return Response(list(users))


@api_view(['GET'])
@permission_classes([IsAdminUser])
def organization_list(request):
    """List all organizations with plan info."""
    orgs = Organization.objects.all()
    serializer = OrganizationSerializer(orgs, many=True)
    return Response(serializer.data)


@api_view(['PATCH'])
@permission_classes([IsAdminUser])
def update_org_plan(request, pk):
    """Update an organization's plan (admin only)."""
    try:
        org = Organization.objects.get(pk=pk)
    except Organization.DoesNotExist:
        return Response(
            {'detail': 'Organization not found.'},
            status=status.HTTP_404_NOT_FOUND
        )

    new_plan = request.data.get('plan')
    if new_plan not in PLAN_LIMITS:
        return Response(
            {'detail': f'Invalid plan. Options: {", ".join(PLAN_LIMITS.keys())}'},
            status=status.HTTP_400_BAD_REQUEST
        )

    org.set_plan(new_plan)
    return Response({
        'detail': f'Plan updated to {new_plan}.',
        'organization': OrganizationSerializer(org).data,
    })


@api_view(['GET', 'PUT', 'DELETE'])
@permission_classes([IsAdminUser])
def organization_detail(request, pk):
    """Retrieve, update or delete an organization."""
    try:
        org = Organization.objects.get(pk=pk)
    except Organization.DoesNotExist:
        return Response({'detail': 'Organization not found.'}, status=status.HTTP_404_NOT_FOUND)

    if request.method == 'GET':
        return Response(OrganizationSerializer(org).data)

    elif request.method == 'PUT':
        from .serializers import AdminOrgUpdateSerializer
        before_plan = org.plan
        serializer = AdminOrgUpdateSerializer(org, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            if 'plan' in request.data and org.plan != before_plan:
                AdminActionLog.objects.create(
                    action='plan_change',
                    performed_by=request.user,
                    performed_by_email=request.user.email,
                    target_org_id=org.id,
                    target_org_name=org.name,
                    before_value=before_plan,
                    after_value=org.plan,
                )
            return Response(OrganizationSerializer(org).data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    elif request.method == 'DELETE':
        org.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


@api_view(['GET', 'PUT', 'DELETE'])
@permission_classes([IsAdminUser])
def user_detail(request, pk):
    """Retrieve, update or delete a user."""
    from accounts.serializers import UserSerializer
    try:
        user = User.objects.get(pk=pk)
    except User.DoesNotExist:
        return Response({'detail': 'User not found.'}, status=status.HTTP_404_NOT_FOUND)

    if request.method == 'GET':
        return Response(UserSerializer(user).data)

    elif request.method == 'PUT':
        from .serializers import AdminUserUpdateSerializer
        if user.pk == request.user.pk and request.data.get('is_active') is False:
            return Response(
                {'detail': 'No puedes bloquear tu propia cuenta.'},
                status=status.HTTP_400_BAD_REQUEST,
            )
        before_role = user.role
        before_org_id = user.organization_id
        before_org_name = str(user.organization) if user.organization else ''
        before_active = user.is_active
        serializer = AdminUserUpdateSerializer(user, data=request.data, partial=True)
        if serializer.is_valid():
            serializer.save()
            if 'role' in request.data and user.role != before_role:
                AdminActionLog.objects.create(
                    action='role_change',
                    performed_by=request.user,
                    performed_by_email=request.user.email,
                    target_user_id=user.id,
                    target_user_email=user.email,
                    before_value=before_role,
                    after_value=user.role,
                )
            if 'organization' in request.data and user.organization_id != before_org_id:
                AdminActionLog.objects.create(
                    action='org_change',
                    performed_by=request.user,
                    performed_by_email=request.user.email,
                    target_user_id=user.id,
                    target_user_email=user.email,
                    before_value=before_org_name,
                    after_value=str(user.organization) if user.organization else '',
                )
            if 'is_active' in request.data and user.is_active != before_active:
                AdminActionLog.objects.create(
                    action='block_change',
                    performed_by=request.user,
                    performed_by_email=request.user.email,
                    target_user_id=user.id,
                    target_user_email=user.email,
                    before_value='active' if before_active else 'blocked',
                    after_value='active' if user.is_active else 'blocked',
                )
            return Response(UserSerializer(user).data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    elif request.method == 'DELETE':
        user.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)


@api_view(['GET'])
@permission_classes([IsAdminUser])
def analytics_organizations(request):
    """Get global ranking and metrics for all organizations."""
    from django.db.models import Count, Q, Max, Sum
    
    orgs = Organization.objects.annotate(
        total_urls=Count('url_assets', distinct=True),
        total_users=Count('users', distinct=True),
        total_scans=Count('url_assets__scans', distinct=True),
        last_scan_date=Max('url_assets__scans__started_at'),
        critical_findings=Count('url_assets__scans__findings', filter=Q(url_assets__scans__findings__severity='CRITICAL'), distinct=True),
        high_findings=Count('url_assets__scans__findings', filter=Q(url_assets__scans__findings__severity='HIGH'), distinct=True),
        medium_findings=Count('url_assets__scans__findings', filter=Q(url_assets__scans__findings__severity='MEDIUM'), distinct=True),
    ).order_by('-critical_findings', '-high_findings')

    data = []
    for org in orgs:
        risk_score = org.critical_findings * 10 + org.high_findings * 5 + org.medium_findings * 2
        
        status = 'Sano'
        if risk_score > 50 or org.critical_findings > 0:
            status = 'Crítico'
        elif risk_score > 20 or org.high_findings > 0:
            status = 'Advertencia'

        data.append({
            'id': org.id,
            'name': org.name,
            'plan': org.plan,
            'total_urls': org.total_urls,
            'total_users': org.total_users,
            'total_scans': org.total_scans,
            'last_scan_date': org.last_scan_date,
            'critical_findings': org.critical_findings,
            'high_findings': org.high_findings,
            'medium_findings': org.medium_findings,
            'risk_score': risk_score,
            'status': status
        })

    return Response(data)


@api_view(['GET'])
@permission_classes([IsAdminUser])
def analytics_organization_detail(request, pk):
    """Get detailed analytics for a specific organization."""
    try:
        org = Organization.objects.get(pk=pk)
    except Organization.DoesNotExist:
        return Response({'detail': 'Organization not found.'}, status=status.HTTP_404_NOT_FOUND)

    # Users — backfill real Clerk emails in place of the placeholder
    # {clerk_id}@clerk.user before serializing (same as user_list).
    from accounts.clerk_api import backfill_real_emails, is_placeholder_email
    org_users = list(org.users.all())
    if any(is_placeholder_email(u.email) for u in org_users):
        backfill_real_emails(org_users)
    users_data = org.users.values('id', 'email', 'role', 'last_login')

    # URLs and their last scan status
    urls = org.url_assets.all()
    urls_data = []
    
    scan_history = []
    total_findings = {'CRITICAL': 0, 'HIGH': 0, 'MEDIUM': 0, 'LOW': 0, 'INFO': 0}

    for url in urls:
        last_scan = url.scans.order_by('-started_at').first()
        
        url_info = {
            'id': url.id,
            'url': url.url,
            'status': url.last_scan_status,
            'last_scan_status': last_scan.status if last_scan else None,
            'last_scan_date': last_scan.started_at if last_scan else None,
            'critical': last_scan.critical_count if last_scan else 0,
            'high': last_scan.high_count if last_scan else 0,
            'medium': last_scan.medium_count if last_scan else 0,
        }
        urls_data.append(url_info)
        
        # Accumulate finding totals from the LAST scan of each URL to avoid counting historical duplicates
        if last_scan:
            total_findings['CRITICAL'] += url_info['critical']
            total_findings['HIGH'] += url_info['high']
            total_findings['MEDIUM'] += url_info['medium']
            total_findings['LOW'] += last_scan.low_count
            
        # Collect recent scan history for charts (limit to last 10 per URL)
        for scan in url.scans.order_by('-started_at')[:10]:
            scan_history.append({
                'id': scan.id,
                'url': url.url,
                'started_at': scan.started_at,
                'status': scan.status,
                'findings_count': scan.findings_count
            })

    # Sort scan history chronologically for chart
    scan_history.sort(key=lambda x: x['started_at'])

    return Response({
        'organization': {
            'id': org.id,
            'name': org.name,
            'plan': org.plan,
        },
        'total_findings': total_findings,
        'users': list(users_data),
        'urls': urls_data,
        'recent_scans': scan_history[-50:], # limit global chart to 50
    })


@api_view(['GET'])
@permission_classes([IsAdminUser])
def audit_log_list(request):
    """List API request audit logs (admin only) — evidence trail, paginated."""
    logs = APIRequestLog.objects.all()

    ip = request.query_params.get('ip')
    if ip:
        logs = logs.filter(ip_address__icontains=ip)

    method = request.query_params.get('method')
    if method:
        logs = logs.filter(method=method.upper())

    search = request.query_params.get('search')
    if search:
        logs = logs.filter(
            Q(path__icontains=search)
            | Q(user_email__icontains=search)
            | Q(organization_name__icontains=search)
        )

    try:
        page = max(int(request.query_params.get('page', 1)), 1)
    except (TypeError, ValueError):
        page = 1
    try:
        page_size = min(max(int(request.query_params.get('page_size', 50)), 1), 200)
    except (TypeError, ValueError):
        page_size = 50

    total = logs.count()
    start = (page - 1) * page_size
    results = logs[start:start + page_size].values(
        'id', 'created_at', 'method', 'path', 'query_string', 'status_code',
        'ip_address', 'user_id', 'user_email', 'organization_id',
        'organization_name', 'user_agent', 'response_time_ms',
    )

    return Response({
        'count': total,
        'page': page,
        'page_size': page_size,
        'results': list(results),
    })


@api_view(['GET'])
@permission_classes([IsAdminUser])
def admin_action_log_list(request):
    """List admin action audit logs (who changed what) — paginated.

    Distinct from audit_log_list (raw HTTP request trail): this is the
    explicit before/after diff for sensitive actions — role changes,
    org moves, plan changes, block/unblock — so an admin can answer
    "who did this and what did it change" without querying the DB.
    """
    logs = AdminActionLog.objects.all()

    action = request.query_params.get('action')
    if action:
        logs = logs.filter(action=action)

    search = request.query_params.get('search')
    if search:
        logs = logs.filter(
            Q(performed_by_email__icontains=search)
            | Q(target_user_email__icontains=search)
            | Q(target_org_name__icontains=search)
        )

    try:
        page = max(int(request.query_params.get('page', 1)), 1)
    except (TypeError, ValueError):
        page = 1
    try:
        page_size = min(max(int(request.query_params.get('page_size', 50)), 1), 200)
    except (TypeError, ValueError):
        page_size = 50

    total = logs.count()
    start = (page - 1) * page_size
    results = logs[start:start + page_size].values(
        'id', 'created_at', 'action', 'performed_by_email',
        'target_user_id', 'target_user_email', 'target_org_id',
        'target_org_name', 'before_value', 'after_value',
    )

    return Response({
        'count': total,
        'page': page,
        'page_size': page_size,
        'results': list(results),
    })

@api_view(['GET'])
@permission_classes([IsAdminUser])
def cyber_map_data(request):
    """Returns recent scans with source and destination IP geolocations."""
    days = request.query_params.get('days', 'general')
    scans_qs = Scan.objects.select_related('created_by', 'url_asset').order_by('-started_at')

    if days == '1':
        scans_qs = scans_qs.filter(started_at__gte=timezone.now() - timedelta(days=1))
    elif days == '7':
        scans_qs = scans_qs.filter(started_at__gte=timezone.now() - timedelta(days=7))
    elif days == '30':
        scans_qs = scans_qs.filter(started_at__gte=timezone.now() - timedelta(days=30))

    recent_scans = scans_qs[:100]
    server_public_ip = get_server_public_ip()

    data = []
    for scan in recent_scans:
        origin_log = APIRequestLog.objects.filter(
            user=scan.created_by
        ).order_by('-created_at').first()

        origin_ip = origin_log.ip_address if origin_log else server_public_ip
        real_origin_ip = server_public_ip if origin_ip in ('127.0.0.1', 'localhost', '0.0.0.0', '::1') else origin_ip
        target_url = scan.url_asset.url
        target_ip = get_target_ip(target_url)

        data.append({
            'id': scan.id,
            'user': scan.created_by.email if scan.created_by else 'Sistema',
            'origin_ip': real_origin_ip,
            'origin_geo': get_geo_info(real_origin_ip, server_public_ip),
            'target_url': target_url,
            'target_ip': target_ip,
            'target_geo': get_geo_info(target_ip, server_public_ip),
            'status': scan.status,
            'timestamp': scan.started_at,
        })

    return Response(data)


# ----------------------------------------------------------------------
# Admin global report (origin/audited countries, top users, top routes)
# ----------------------------------------------------------------------

def _get_scans_for_report(days='general', limit=200):
    """Return recent scans filtered by the same 'days' param used by cyber-map."""
    scans_qs = Scan.objects.select_related('created_by', 'url_asset').order_by('-started_at')

    if days == '1':
        scans_qs = scans_qs.filter(started_at__gte=timezone.now() - timedelta(days=1))
    elif days == '7':
        scans_qs = scans_qs.filter(started_at__gte=timezone.now() - timedelta(days=7))
    elif days == '30':
        scans_qs = scans_qs.filter(started_at__gte=timezone.now() - timedelta(days=30))

    return list(scans_qs[:limit])


def _collect_admin_report_data(days='general'):
    """Aggregate data for the admin global report."""
    scans = _get_scans_for_report(days)
    server_public_ip = get_server_public_ip()

    status_counts = Counter()
    user_counts = Counter()
    target_counts = Counter()
    origin_country_counts = Counter()
    target_country_counts = Counter()

    for scan in scans:
        status_counts[scan.status] += 1

        user_email = scan.created_by.email if scan.created_by else 'Sistema'
        user_counts[user_email] += 1

        target_url = scan.url_asset.url if scan.url_asset else 'N/A'
        target_counts[target_url] += 1

        # Origin country from the user's most recent APIRequestLog
        origin_log = APIRequestLog.objects.filter(
            user=scan.created_by
        ).order_by('-created_at').first()
        origin_ip = origin_log.ip_address if origin_log else server_public_ip
        real_origin_ip = server_public_ip if origin_ip in ('127.0.0.1', 'localhost', '0.0.0.0', '::1') else origin_ip
        origin_geo = get_geo_info(real_origin_ip, server_public_ip)
        origin_country_counts[origin_geo.get('country', 'Desconocido')] += 1

        # Target country from DNS resolution of the scanned URL
        target_ip = get_target_ip(target_url)
        target_geo = get_geo_info(target_ip, server_public_ip)
        target_country_counts[target_geo.get('country', 'Desconocido')] += 1

    return {
        'period': days,
        'generated_at': timezone.now(),
        'total_scans': len(scans),
        'status_counts': dict(status_counts),
        'top_origin_countries': [
            {'country': country, 'count': count}
            for country, count in origin_country_counts.most_common(10)
        ],
        'top_target_countries': [
            {'country': country, 'count': count}
            for country, count in target_country_counts.most_common(10)
        ],
        'top_users': [
            {'user': user, 'count': count}
            for user, count in user_counts.most_common(10)
        ],
        'top_targets': [
            {'url': url, 'count': count}
            for url, count in target_counts.most_common(10)
        ],
    }


@api_view(['GET'])
@permission_classes([IsAdminUser])
def admin_report_summary(request):
    """Return aggregated admin report data as JSON."""
    days = request.query_params.get('days', 'general')
    data = _collect_admin_report_data(days)
    data['generated_at'] = data['generated_at'].isoformat()
    return Response(data)


def _build_admin_excel(data, days):
    """Build an Excel workbook for the admin global report."""
    wb = Workbook()

    header_font = Font(name='Calibri', size=11, bold=True, color='FFFFFFFF')
    header_fill = PatternFill(start_color='FF1E293B', end_color='FF1E293B', fill_type='solid')
    thin_side = Side(border_style='thin', color='FFCBD5E1')
    border = Border(left=thin_side, right=thin_side, top=thin_side, bottom=thin_side)
    cell_align = Alignment(wrap_text=True, vertical='top')

    period_label = {
        'general': 'Todos los tiempos (últimos 200 escaneos)',
        '1': 'Últimas 24 horas',
        '7': 'Últimos 7 días',
        '30': 'Últimos 30 días',
    }.get(days, days)
    generated_str = data['generated_at'].strftime('%d/%m/%Y %H:%M')
    subtitle = f'Período: {period_label} · Generado: {generated_str}'

    def _auto_width(ws, skip_rows=0):
        for col in ws.columns:
            max_len = 0
            for cell in col:
                if cell.row <= skip_rows:
                    continue
                if cell.value is not None:
                    cell.border = border
                    cell.alignment = cell_align
                    max_len = max(max_len, len(str(cell.value)))
            ws.column_dimensions[get_column_letter(col[0].column)].width = min(max_len + 2, 60)

    def _add_table(ws, title, headers, rows, tab_color):
        _setup_sheet_header(ws, title, subtitle, col_span=len(headers))
        header_row = ws.max_row + 1
        ws.append(headers)
        for cell in ws[header_row]:
            cell.font = header_font
            cell.fill = header_fill
            cell.border = border
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for row in rows:
            ws.append(row)
        _auto_width(ws, skip_rows=1)
        _finalize_table_sheet(ws, header_row, tab_color=tab_color, zebra=True)

    # Sheet 1: Resumen
    ws1 = wb.active
    ws1.title = 'Resumen'
    _setup_sheet_header(ws1, 'Reporte Global de Auditoría', subtitle, col_span=2)
    header_row = ws1.max_row + 1
    ws1.append(['Campo', 'Valor'])
    for cell in ws1[header_row]:
        cell.font = header_font
        cell.fill = header_fill
        cell.border = border
        cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)

    summary_rows = [
        ['Período', period_label],
        ['Total de escaneos', data['total_scans']],
        ['Completados', data['status_counts'].get('completed', 0)],
        ['En ejecución', data['status_counts'].get('running', 0)],
        ['Pendientes', data['status_counts'].get('pending', 0)],
        ['Con error', data['status_counts'].get('error', 0)],
    ]
    for row in summary_rows:
        ws1.append(row)
    _auto_width(ws1, skip_rows=1)
    _finalize_table_sheet(ws1, header_row, tab_color='2563EB', zebra=True)

    # Sheet 2: Países de origen
    ws2 = wb.create_sheet('Países Origen')
    _add_table(
        ws2,
        'Países de Origen (más peticiones)',
        ['País', 'Escaneos'],
        [[item['country'], item['count']] for item in data['top_origin_countries']],
        '0891B2',
    )

    # Sheet 3: Países auditados
    ws3 = wb.create_sheet('Países Auditados')
    _add_table(
        ws3,
        'Países Auditados (más escaneados)',
        ['País', 'Escaneos'],
        [[item['country'], item['count']] for item in data['top_target_countries']],
        '7C3AED',
    )

    # Sheet 4: Usuarios
    ws4 = wb.create_sheet('Usuarios')
    _add_table(
        ws4,
        'Usuarios con más peticiones',
        ['Usuario', 'Escaneos'],
        [[item['user'], item['count']] for item in data['top_users']],
        '059669',
    )

    # Sheet 5: Rutas
    ws5 = wb.create_sheet('Rutas')
    _add_table(
        ws5,
        'Rutas más auditadas',
        ['URL', 'Escaneos'],
        [[item['url'], item['count']] for item in data['top_targets']],
        'DC2626',
    )

    buffer = BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer


def _build_admin_pdf(data, days):
    """Build a PDF for the admin global report using ReportLab."""
    if not REPORTLAB_AVAILABLE:
        raise RuntimeError('ReportLab no está disponible')

    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=LETTER,
        leftMargin=0.5 * inch,
        rightMargin=0.5 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
        title='Reporte Global de Auditoria',
    )

    styles = getSampleStyleSheet()
    title_style = ParagraphStyle(
        'AdminTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=18,
        textColor=rl_colors.HexColor('#0f172a'),
        leading=22,
        spaceAfter=4,
    )
    meta_style = ParagraphStyle(
        'AdminMeta',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        textColor=rl_colors.HexColor('#475569'),
        leading=12,
        spaceAfter=12,
    )
    section_style = ParagraphStyle(
        'AdminSection',
        parent=styles['Heading3'],
        fontName='Helvetica-Bold',
        fontSize=11,
        textColor=rl_colors.HexColor('#0f172a'),
        spaceBefore=10,
        spaceAfter=4,
    )
    body_style = ParagraphStyle(
        'AdminBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=rl_colors.HexColor('#1f2937'),
    )
    header_style = ParagraphStyle(
        'AdminTableHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        textColor=rl_colors.white,
        leading=11,
    )

    period_label = {
        'general': 'Todos los tiempos (últimos 200 escaneos)',
        '1': 'Últimas 24 horas',
        '7': 'Últimos 7 días',
        '30': 'Últimos 30 días',
    }.get(days, days)
    generated_str = data['generated_at'].strftime('%d/%m/%Y %H:%M')

    story = []

    if LOGO_AVAILABLE:
        try:
            img = RLImage(LOGO_PATH, width=160, height=160 * (120 / 420))
            img.hAlign = 'LEFT'
            story.append(img)
            story.append(Spacer(1, 6))
        except Exception:
            pass

    story.append(Paragraph('Reporte Global de Auditoría', title_style))
    story.append(Paragraph(
        f'Período: {period_label} · Total escaneos: {data["total_scans"]} · Generado: {generated_str}',
        meta_style,
    ))

    # Summary stats
    story.append(Paragraph('Resumen de Estado', section_style))
    stat_headers = ['Completados', 'En ejecución', 'Pendientes', 'Con error']
    stat_values = [
        data['status_counts'].get('completed', 0),
        data['status_counts'].get('running', 0),
        data['status_counts'].get('pending', 0),
        data['status_counts'].get('error', 0),
    ]
    stats_data = [
        [Paragraph(h, header_style) for h in stat_headers],
        [Paragraph(str(v), body_style) for v in stat_values],
    ]
    page_width = 7.5 * inch
    stats_tbl = Table(stats_data, repeatRows=1, colWidths=[page_width / 4] * 4)
    stats_tbl.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), rl_colors.HexColor('#1e293b')),
        ('GRID', (0, 0), (-1, -1), 0.5, rl_colors.HexColor('#cbd5e1')),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(stats_tbl)

    def _clean(text):
        return (str(text or '')).replace('&', '&amp;').replace('<', '&lt;').replace('>', '&gt;')

    def _add_ranking_table(title, headers, items, key_label, key_key, count_key):
        story.append(Paragraph(title, section_style))
        rows = [[Paragraph(h, header_style) for h in headers]]
        for item in items:
            rows.append([
                Paragraph(_clean(item[key_label]), body_style),
                Paragraph(str(item[count_key]), body_style),
            ])
        if not items:
            rows.append([Paragraph('Sin datos', body_style), Paragraph('-', body_style)])
        col_widths = [page_width - 1.2 * inch, 1.2 * inch]
        tbl = Table(rows, repeatRows=1, colWidths=col_widths)
        tbl.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), rl_colors.HexColor('#1e293b')),
            ('GRID', (0, 0), (-1, -1), 0.3, rl_colors.HexColor('#cbd5e1')),
            ('VALIGN', (0, 0), (-1, -1), 'TOP'),
            ('LEFTPADDING', (0, 0), (-1, -1), 4),
            ('RIGHTPADDING', (0, 0), (-1, -1), 4),
            ('TOPPADDING', (0, 0), (-1, -1), 4),
            ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(tbl)
        story.append(Spacer(1, 6))

    _add_ranking_table(
        'Países de Origen',
        ['País', 'Escaneos'],
        data['top_origin_countries'],
        'country',
        'country',
        'count',
    )
    _add_ranking_table(
        'Países Auditados',
        ['País', 'Escaneos'],
        data['top_target_countries'],
        'country',
        'country',
        'count',
    )
    _add_ranking_table(
        'Usuarios con más peticiones',
        ['Usuario', 'Escaneos'],
        data['top_users'],
        'user',
        'user',
        'count',
    )
    _add_ranking_table(
        'Rutas más auditadas',
        ['URL', 'Escaneos'],
        data['top_targets'],
        'url',
        'url',
        'count',
    )

    story.append(Paragraph(
        'Generado por Auditoría Web Automatizada para PYMES',
        ParagraphStyle(
            'AdminFooter',
            parent=styles['Normal'],
            alignment=1,
            fontSize=8,
            textColor=rl_colors.HexColor('#64748b'),
        ),
    ))

    doc.build(story)
    buffer.seek(0)
    return buffer


@api_view(['GET'])
@permission_classes([IsAdminUser])
def admin_report_export(request):
    """Download the admin global report as Excel or PDF."""
    days = request.query_params.get('days', 'general')
    fmt = (request.query_params.get('format') or 'excel').lower()

    data = _collect_admin_report_data(days)

    safe_period = {'general': 'todos', '1': '24h', '7': '7d', '30': '30d'}.get(days, days)
    timestamp = data['generated_at'].strftime('%Y%m%d_%H%M')

    if fmt in ('excel', 'xlsx'):
        buffer = _build_admin_excel(data, days)
        response = HttpResponse(
            buffer.getvalue(),
            content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',
        )
        response['Content-Disposition'] = (
            f'attachment; filename="reporte_global_{safe_period}_{timestamp}.xlsx"'
        )
        return response

    if fmt == 'pdf':
        if not REPORTLAB_AVAILABLE:
            return Response(
                {'detail': 'ReportLab no está disponible para generar PDF.'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )
        try:
            buffer = _build_admin_pdf(data, days)
            response = HttpResponse(buffer.getvalue(), content_type='application/pdf')
            response['Content-Disposition'] = (
                f'attachment; filename="reporte_global_{safe_period}_{timestamp}.pdf"'
            )
            return response
        except Exception as exc:
            return Response(
                {'detail': f'Error generando PDF: {exc}'},
                status=status.HTTP_500_INTERNAL_SERVER_ERROR,
            )

    return Response(
        {'detail': 'Formato no soportado. Usa format=excel o format=pdf.'},
        status=status.HTTP_400_BAD_REQUEST,
    )
